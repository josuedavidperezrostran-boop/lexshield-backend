from __future__ import annotations

import io

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}


def extract_text(filename: str, content_type: str | None, payload: bytes) -> str:
    kind = _detect_kind(filename, content_type)
    if kind == "pdf":
        text = _from_pdf(payload)
    else:
        text = _from_docx(payload)

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < 40:
        raise HTTPException(
            status_code=400,
            detail="Could not extract readable contract text from this file.",
        )
    return cleaned


def _detect_kind(filename: str, content_type: str | None) -> str:
    lowered = filename.lower()
    if lowered.endswith(".doc") and not lowered.endswith(".docx"):
        raise HTTPException(
            status_code=415,
            detail="Legacy .doc files are not supported. Save as .docx or .pdf.",
        )
    if lowered.endswith(".pdf") or content_type == "application/pdf":
        return "pdf"
    if lowered.endswith(".docx") or content_type in SUPPORTED_TYPES:
        return "docx"
    raise HTTPException(
        status_code=415,
        detail="Only .pdf and .docx contracts are accepted. Files never leave memory.",
    )


def _from_pdf(payload: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(payload))
        pages = [
            f"[[PÁGINA {index}]]\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages, start=1)
        ]
    except Exception as exc:  # noqa: BLE001 — surface a clean client error
        raise HTTPException(status_code=400, detail="Unable to parse this PDF.") from exc
    return "\n\n".join(pages)


def _from_docx(payload: bytes) -> str:
    try:
        document = Document(io.BytesIO(payload))
        paragraphs = [p.text for p in document.paragraphs]
        tables = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    except Exception as extra:  # noqa: BLE001
        raise HTTPException(status_code=400, detail="Unable to parse this DOCX.") from extra
    return "[[PÁGINA 1]]\n" + "\n".join(paragraphs + tables)
